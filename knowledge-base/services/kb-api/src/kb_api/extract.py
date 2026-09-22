"""Extracao do documento canonico (ING-03 e ING-04).

O canonico e Markdown limpo: o texto que a busca ve, versionado, derivado do
bruto que fica imutavel no object store. Nenhuma das tres ferramentas de mercado
avaliadas produz esse artefato -- e por isso ele e o centro desta implementacao.

**Mais de uma tecnica disponivel, comparavel** e requisito (ING-04), nao enfeite:

1. `docling` -- layout-aware, tabelas em Markdown, o melhor de mercado para PDF
   e o que a propria lista de requisitos cita como referencia. E pesado (traz
   torch), entao entra por import tardio: se nao estiver na imagem, o pipeline
   continua com o proximo extrator em vez de morrer.
2. `pymupdf` / `python-docx` / `python-pptx` / `openpyxl` -- rapido, sem modelo
   de ML, cobre o mesmo conjunto de formatos com menos fidelidade de layout.
3. texto puro -- ultimo recurso para md/txt/csv/json.

Qual foi usado fica gravado em `document.extractor`, para a comparacao entre
tecnicas ser possivel depois (FUN-04).
"""

from __future__ import annotations

import io
import logging
import re
import subprocess
import time
from dataclasses import dataclass, field
from pathlib import Path

from . import okf
from .config import settings

log = logging.getLogger(__name__)

# Ordem de preferencia por extensao. O primeiro que conseguir, ganha.
DOCLING_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".md", ".adoc"}
PLAIN_EXTENSIONS = {".md", ".txt", ".csv", ".json", ".properties", ".vtt"}

SUPPORTED = DOCLING_EXTENSIONS | PLAIN_EXTENSIONS | {".doc", ".xls", ".ppt", ".eml"}


class ExtractionError(RuntimeError):
    pass


@dataclass
class Figure:
    """Uma imagem do documento, com o texto que o OCR leu dentro dela.

    Existe porque imagem em documento corporativo raramente e decoracao: e o
    print da tela do sistema, o fluxograma do processo, a tabela que alguem
    colou como figura. Sem OCR, esse conteudo simplesmente nao existe para a
    busca -- o documento aparece indexado e a resposta certa esta numa imagem
    que ninguem leu.
    """

    ref: str  # identificador estavel dentro do documento (fig-1, fig-2, ...)
    page: int | None
    data: bytes  # PNG
    caption: str = ""
    ocr_text: str = ""


@dataclass
class Extracted:
    markdown: str
    extractor: str
    title: str = ""
    # Numero de paginas, quando o formato tem paginacao. Serve para a UI abrir o
    # original na pagina certa.
    pages: int = 0
    figures: list[Figure] = field(default_factory=list)


def _clean(text: str) -> str:
    """Normaliza espaco em branco sem reescrever o conteudo.

    O requisito e explicito: extrair SEM reescrever (ING-03). Entao aqui so cai
    espaco redundante e linha vazia em excesso -- nada de resumir, reordenar ou
    "melhorar" o texto.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# Marca de pagina no canonico. E o unico jeito de, mais tarde, dizer em QUE
# pagina do original um trecho estava -- a UI usa isso para abrir o PDF na
# pagina certa em vez de largar o arquivo inteiro na mao de quem pergunta.
PAGE_MARK = "<!-- pagina {page} -->"
PAGE_MARK_RE = re.compile(r"<!--\s*pagina\s+(\d+)\s*-->")

# Referencia de figura no canonico. Fica em linha propria, com a legenda e o
# texto lido por OCR logo abaixo, para o chunk que cair ali carregar as duas
# coisas: a evidencia textual E o ponteiro para a imagem.
FIGURE_MARK = "<!-- figura {ref} -->"
FIGURE_MARK_RE = re.compile(r"<!--\s*figura\s+([a-zA-Z0-9_-]+)\s*-->")


def page_of_offset(markdown: str, offset: int) -> int | None:
    """Pagina do original correspondente a um ponto do canonico.

    A ultima marca de pagina ANTES do offset e a resposta. Documento sem marca
    (docx, txt) devolve None -- e melhor nao dizer nada do que apontar a pagina
    errada.
    """
    page = None
    for match in PAGE_MARK_RE.finditer(markdown):
        if match.start() > offset:
            break
        page = int(match.group(1))
    return page


# Linha que existe por estrutura, nao por conteudo: nunca deve virar titulo.
_NAO_E_TITULO = re.compile(
    r"^(<!--.*|!\[.*\]\(.*\)|_[^_]*_|\*\*[^*]*\*\*)$"
)

# Frontmatter YAML no topo do arquivo. Nao e conteudo, e nunca foi -- mas ate
# aqui o `_first_heading` lia a primeira linha dele, que e `---`, e era isso que
# ia para a tela e para a citacao do agente como titulo do documento. Aparece em
# todo conceito OKF e em qualquer Markdown de gerador estatico.
_FRONTMATTER = re.compile(r"\A---[ \t]*\r?\n.*?\r?\n---[ \t]*(?:\r?\n|\Z)", re.DOTALL)


def _first_heading(markdown: str, fallback: str) -> str:
    """Titulo do documento: o primeiro cabecalho ou a primeira linha de texto.

    As guardas nao sao paranoia -- cada uma corrige um titulo que ja apareceu
    errado na tela e, pior, como fonte citada pelo agente:

    * comentario HTML: a marca de pagina virava "<!-- pagina 1 -->";
    * imagem Markdown e legenda em negrito: documento que comeca por figura
      virava "_imagem sem texto legivel_" ou o texto da legenda solto;
    * frontmatter YAML: arquivo que comeca com `---` virava um titulo "---". E
      o caso de todo conceito OKF, e de qualquer Markdown de gerador estatico.
    """
    corte = _FRONTMATTER.match(markdown)
    for line in markdown[corte.end():].splitlines() if corte else markdown.splitlines():
        stripped = line.strip()
        if not stripped or _NAO_E_TITULO.match(stripped):
            continue
        if stripped.startswith("#"):
            titulo = stripped.lstrip("#").strip()
            if titulo:
                return titulo
            continue
        return stripped[:180]
    return fallback


# ── 1. docling ─────────────────────────────────────────────────────────────


_CONVERTER = None


def _docling_converter():
    """Converter do docling, criado UMA vez por processo.

    Instanciar `DocumentConverter()` por arquivo recarrega os modelos de layout
    a cada documento -- o que, com 44 arquivos, e a diferenca entre uma ingestao
    de minutos e uma de dezenas de minutos, alem de picos de memoria repetidos.

    OCR: ligado, mas NAO forcado. Com `force_full_page_ocr=False` o docling so
    chama o OCR onde nao ha camada de texto -- imagem, print de tela, pagina
    digitalizada. PDF digital continua saindo pelo caminho rapido, entao o custo
    aparece apenas onde ha ganho de verdade.

    O motor e o tesseract por CLI, nao easyocr. Nao e preferencia estetica: o
    easyocr carrega mais um modelo no MESMO processo do uvicorn, e este pod ja
    foi OOMKilled por carregar modelo duas vezes. O tesseract roda como
    subprocesso, devolve o texto e morre -- a memoria volta.
    """
    global _CONVERTER
    if _CONVERTER is not None:
        return _CONVERTER

    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    options = PdfPipelineOptions()
    options.do_table_structure = True
    options.do_ocr = settings.ocr_enabled
    if settings.ocr_enabled:
        options.ocr_options = TesseractCliOcrOptions(
            lang=settings.ocr_languages,
            force_full_page_ocr=settings.ocr_force_full_page,
            # psm explicito desliga a deteccao de orientacao (OSD) do tesseract.
            # Sem isso, cada regiao sem texto suficiente produzia
            # `OSD failed ... Too few characters` em nivel ERROR, e um PDF de 30
            # paginas enchia o log com dezenas de erros que nao eram erros --
            # atrapalhando achar as falhas de verdade. E cada OSD e uma chamada
            # extra ao binario, entao tirar tambem economiza tempo.
            # 3 = segmentacao automatica de pagina, sem OSD.
            psm=settings.ocr_psm,
        )
    # As imagens das figuras sao geradas para poderem ser guardadas e
    # referenciadas. `images_scale` acima de 1 melhora o OCR da figura; 2 e o
    # ponto em que o texto de print de tela passa a ser lido de forma confiavel
    # sem o PNG ficar grande demais para guardar por documento.
    options.generate_picture_images = settings.figures_enabled
    options.images_scale = settings.figure_scale
    _CONVERTER = DocumentConverter(
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
    )
    log.info(
        "docling pronto (ocr=%s, idiomas=%s, figuras=%s)",
        settings.ocr_enabled, ",".join(settings.ocr_languages), settings.figures_enabled,
    )
    return _CONVERTER


def _page_count(document) -> int:
    try:
        pages = getattr(document, "pages", None)
        return len(pages) if pages else 0
    except Exception:  # noqa: BLE001
        return 0


def _docling_markdown(document) -> tuple[str, int]:
    """Markdown do documento COM marca de pagina.

    O `export_to_markdown()` inteiro nao diz onde uma pagina termina, e sem isso
    nao ha como levar quem pergunta ate a pagina certa do original -- so ate o
    arquivo. Exportar pagina por pagina resolve, e o custo e uma passada a mais
    sobre uma arvore que ja esta em memoria.

    `escape_html=False` de proposito: com o default, `GH & VOCE` era gravado no
    canonico como `GH &amp; VOCE`, e a entidade aparecia crua no titulo e na
    evidencia devolvida ao agente. O canonico e texto, nao HTML.
    """
    total = _page_count(document)
    if total <= 0:
        return _clean(document.export_to_markdown(escape_html=False)), 0

    parts: list[str] = []
    for page in range(1, total + 1):
        try:
            body = document.export_to_markdown(page_no=page, escape_html=False)
        except Exception as exc:  # noqa: BLE001 - versao sem page_no
            log.info("docling sem exportacao por pagina (%s); canonico sem marca", exc)
            return _clean(document.export_to_markdown(escape_html=False)), total
        body = body.strip()
        if body:
            parts.append(f"{PAGE_MARK.format(page=page)}\n{body}")
    return _clean("\n\n".join(parts)), total


def _ocr_image(data: bytes) -> str:
    """Texto de uma imagem isolada, pelo tesseract CLI.

    Roda por subprocesso e le do stdout: sem binding Python, sem modelo carregado
    no processo do servidor. Falha aqui nunca derruba a ingestao -- figura sem
    OCR e uma figura sem texto, nao um documento perdido.
    """
    if not settings.ocr_enabled:
        return ""
    try:
        result = subprocess.run(
            ["tesseract", "stdin", "stdout", "-l", "+".join(settings.ocr_languages)],
            input=data, capture_output=True, timeout=settings.ocr_timeout_seconds,
        )
    except FileNotFoundError:
        log.warning("tesseract nao esta na imagem; figura fica sem OCR")
        return ""
    except subprocess.TimeoutExpired:
        log.warning("OCR da figura passou de %ss", settings.ocr_timeout_seconds)
        return ""
    if result.returncode != 0:
        log.info("tesseract saiu com %s: %s", result.returncode, result.stderr[:200])
        return ""
    return _clean(result.stdout.decode("utf-8", "replace"))


def _docling_figures(document) -> list[Figure]:
    """Figuras do documento, cada uma com legenda e texto lido por OCR."""
    if not settings.figures_enabled:
        return []
    try:
        from docling_core.types.doc import PictureItem
    except Exception:  # noqa: BLE001
        return []

    figures: list[Figure] = []
    total_bytes = 0
    for item, _level in document.iterate_items():
        if not isinstance(item, PictureItem):
            continue
        if len(figures) >= settings.max_figures_per_document:
            log.info("documento passou de %s figuras; o resto fica de fora",
                     settings.max_figures_per_document)
            break
        # Teto do ACUMULADO, nao so de cada figura. As imagens ficam todas em
        # memoria ate a ingestao gravar, e foi exatamente isso que matou o pod
        # por OOM num PDF de 50 paginas cheio de print de tela: cada figura
        # cabia no limite individual, a soma nao.
        if total_bytes >= settings.max_figures_total_bytes:
            log.info(
                "documento passou de %s bytes em figuras; o resto fica de fora",
                settings.max_figures_total_bytes,
            )
            break
        try:
            image = item.get_image(document)
        except Exception:  # noqa: BLE001
            image = None
        if image is None:
            continue
        buffer = io.BytesIO()
        try:
            image.save(buffer, format="PNG")
        except Exception:  # noqa: BLE001
            continue
        finally:
            # Solta o bitmap descomprimido assim que o PNG existe. Uma pagina A4
            # a 2x sao ~8 MB em RGB, e manter dezenas delas vivas e a diferenca
            # entre caber e nao caber.
            image.close()
            del image
        data = buffer.getvalue()
        if len(data) > settings.max_figure_bytes:
            log.info("figura de %s bytes acima do teto; nao sera guardada", len(data))
            continue
        total_bytes += len(data)

        page = None
        provenance = getattr(item, "prov", None) or []
        if provenance:
            page = getattr(provenance[0], "page_no", None)
        caption = ""
        try:
            caption = (item.caption_text(document) or "").strip()
        except Exception:  # noqa: BLE001
            caption = ""

        ref = f"fig-{len(figures) + 1}"
        figures.append(
            Figure(ref=ref, page=page, data=data, caption=caption, ocr_text=_ocr_image(data))
        )
    return figures


def _attach_figures(markdown: str, figures: list[Figure]) -> str:
    """Troca os `<!-- image -->` do docling por referencia de figura + OCR.

    A ordem dos placeholders no Markdown segue a ordem de `iterate_items()`,
    entao o n-esimo placeholder e a n-esima figura. Se as contas nao fecharem
    (placeholder a mais ou a menos), as figuras que sobraram vao para o fim do
    documento: perder a posicao exata e ruim, perder o TEXTO da figura e pior.
    """
    if not figures:
        return markdown

    restantes = list(figures)

    def trocar(_match: re.Match[str]) -> str:
        if not restantes:
            return ""
        figura = restantes.pop(0)
        return _figure_block(figura)

    markdown = re.sub(r"<!--\s*image\s*-->", trocar, markdown)
    if restantes:
        markdown = markdown.rstrip() + "\n\n" + "\n\n".join(
            _figure_block(figura) for figura in restantes
        )
    return _clean(markdown)


def _figure_block(figura: Figure) -> str:
    linhas = [FIGURE_MARK.format(ref=figura.ref)]
    if figura.caption:
        linhas.append(f"**{figura.caption}**")
    if figura.ocr_text:
        # O texto da figura entra no canonico como texto normal, nao como
        # citacao: e ele que a busca lexical e o embedding veem, e um bloco de
        # citacao nao mudaria nada para o indice mas atrapalharia a leitura.
        linhas.append(figura.ocr_text)
    # Figura sem legenda e sem OCR nao ganha texto NENHUM. A versao anterior
    # escrevia "_imagem sem texto legivel_" ali, e isso custou duas coisas: era
    # texto sem conteudo entrando no indice, e em documento que comeca por
    # figura virava o TITULO do documento -- que e o que o agente cita como
    # fonte. A marca sozinha ja carrega a posicao, e a UI diz na galeria quando
    # o OCR nao leu nada.
    return "\n".join(linhas)


def _docling(data: bytes, filename: str) -> Extracted | None:
    try:
        from docling.datamodel.base_models import DocumentStream
    except Exception as exc:  # noqa: BLE001 - ausencia da lib e caso esperado
        log.info("docling indisponivel (%s); usando o extrator seguinte", exc)
        return None

    try:
        converter = _docling_converter()
        stream = DocumentStream(name=filename, stream=io.BytesIO(data))
        result = converter.convert(stream)
        markdown, pages = _docling_markdown(result.document)
        figures = _docling_figures(result.document)
        markdown = _attach_figures(markdown, figures)
    except Exception as exc:  # noqa: BLE001 - arquivo protegido, corrompido etc.
        log.warning("docling falhou em %s: %s", filename, exc)
        return None

    if not markdown:
        return None
    return Extracted(markdown=markdown, extractor="docling", pages=pages, figures=figures)


# ── 2. bibliotecas por formato ─────────────────────────────────────────────


# OCR NO FALLBACK TAMBEM, E NAO SO NO DOCLING.
#
# O docling e o extrator primario de pdf, docx, pptx e xlsx, e e ele que lia as
# figuras. Quando ele falha, a cadeia cai para `pymupdf` / `python-docx` -- que
# produziam markdown SEM figura e SEM OCR. O documento entrava na base
# buscavel, com metade do conteudo faltando, e nada dizia.
#
# Num acervo de manuais isso e grave: o passo a passo mora no print de tela. Um
# manual sem o texto das imagens responde "onde eu clico?" com silencio.
#
# O teto por documento e o mesmo do docling (`max_figures_per_document`), e o
# OCR de cada imagem ja tem timeout proprio -- fallback nao pode virar caminho
# de ingestao infinita.


# Orcamento TOTAL de OCR no fallback, em segundos.
#
# Os tetos que ja existiam se multiplicam: 60 figuras por documento x 120 s de
# timeout por imagem = duas horas de pior caso. No docling isso nao aparece
# porque ele decide sozinho onde vale chamar o OCR e roda em lote; aqui cada
# imagem e um subprocesso de tesseract, em serie.
#
# Medido do jeito mais direto possivel: a primeira versao sem orcamento fez o
# teste ser morto esperando, num PDF de 1,4 MB.
#
# 180 s e a ordem de grandeza da ingestao normal (mediana de 26 s, pior medido
# de 84 s). Passar disso num caminho que ja e o plano B nao vale: melhor um
# documento com as primeiras figuras lidas e as demais sem texto do que uma
# ingestao que parece travada.
ORCAMENTO_OCR_FALLBACK = 180.0


def _figuras_de_imagens(imagens: list[tuple[bytes, int]]) -> list[Figure]:
    """Figuras a partir de bytes crus, com OCR. Sem legenda: o fallback nao a tem.

    Para de chamar o OCR quando o orcamento acaba, mas continua registrando a
    figura: a imagem fica na galeria e o marcador fica no canonico, so sem o
    texto. Perder a posicao e ruim, perder a figura inteira e pior.
    """
    figuras: list[Figure] = []
    comeco = time.monotonic()
    sem_orcamento = 0
    for dados, pagina in imagens[: settings.max_figures_per_document]:
        dentro_do_orcamento = (time.monotonic() - comeco) < ORCAMENTO_OCR_FALLBACK
        if not dentro_do_orcamento:
            sem_orcamento += 1
        figuras.append(
            Figure(
                ref=f"fig-{len(figuras) + 1}",
                page=pagina,
                data=dados,
                caption="",
                ocr_text=_ocr_image(dados) if dentro_do_orcamento else "",
            )
        )
    if sem_orcamento:
        log.warning(
            "orcamento de OCR do fallback (%ss) acabou: %s figura(s) ficaram sem texto",
            ORCAMENTO_OCR_FALLBACK, sem_orcamento,
        )
    return figuras


def _pymupdf(data: bytes, filename: str) -> Extracted | None:
    try:
        import fitz  # PyMuPDF
    except Exception:  # noqa: BLE001
        return None
    try:
        imagens: list[tuple[bytes, int]] = []
        with fitz.open(stream=data, filetype="pdf") as doc:
            pages = []
            for number, page in enumerate(doc, start=1):
                text = page.get_text("text")
                if text.strip():
                    # A marca de pagina preserva a rastreabilidade da fonte.
                    pages.append(f"<!-- pagina {number} -->\n{text}")
                if settings.figures_enabled and len(imagens) < settings.max_figures_per_document:
                    for info in page.get_images(full=True):
                        try:
                            bruto = doc.extract_image(info[0])
                        except Exception:  # noqa: BLE001 - imagem quebrada e so uma figura
                            continue
                        imagens.append((bruto["image"], number))
                        if len(imagens) >= settings.max_figures_per_document:
                            break
        figuras = _figuras_de_imagens(imagens)
        markdown = _attach_figures(_clean("\n\n".join(pages)), figuras)
    except Exception as exc:  # noqa: BLE001
        log.warning("pymupdf falhou em %s: %s", filename, exc)
        return None
    return (
        Extracted(markdown=markdown, extractor="pymupdf", figures=figuras)
        if markdown else None
    )


def _docx(data: bytes, filename: str) -> Extracted | None:
    try:
        import docx  # python-docx
    except Exception:  # noqa: BLE001
        return None
    try:
        document = docx.Document(io.BytesIO(data))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "2"
                blocks.append(f"{'#' * min(int(level), 6)} {text}")
            else:
                blocks.append(text)
        for table in document.tables:
            rows = [
                "| " + " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells) + " |"
                for row in table.rows
            ]
            if len(rows) >= 2:
                header, *body = rows
                sep = "|" + "|".join([" --- "] * len(table.rows[0].cells)) + "|"
                blocks.append("\n".join([header, sep, *body]))
        # As imagens do .docx sao partes do pacote, alcancaveis pelas relacoes
        # do documento. Nao ha posicao confiavel aqui (o `python-docx` nao
        # expoe a ordem no corpo), entao elas vao para o fim -- perder a
        # posicao e ruim, perder o TEXTO e pior, que e a mesma regra do
        # `_attach_figures`.
        imagens: list[tuple[bytes, int]] = []
        if settings.figures_enabled:
            for relacao in document.part.rels.values():
                if "image" not in relacao.reltype:
                    continue
                try:
                    imagens.append((relacao.target_part.blob, 0))
                except Exception:  # noqa: BLE001
                    continue
                if len(imagens) >= settings.max_figures_per_document:
                    break
        figuras = _figuras_de_imagens(imagens)
        markdown = _attach_figures(_clean("\n\n".join(blocks)), figuras)
    except Exception as exc:  # noqa: BLE001
        log.warning("python-docx falhou em %s: %s", filename, exc)
        return None
    return (
        Extracted(markdown=markdown, extractor="python-docx", figures=figuras)
        if markdown else None
    )


def _pptx(data: bytes, filename: str) -> Extracted | None:
    try:
        from pptx import Presentation
    except Exception:  # noqa: BLE001
        return None
    try:
        presentation = Presentation(io.BytesIO(data))
        blocks = []
        for number, slide in enumerate(presentation.slides, start=1):
            lines = [
                shape.text.strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            if lines:
                blocks.append(f"## Slide {number}\n\n" + "\n\n".join(lines))
        markdown = _clean("\n\n".join(blocks))
    except Exception as exc:  # noqa: BLE001
        log.warning("python-pptx falhou em %s: %s", filename, exc)
        return None
    return Extracted(markdown=markdown, extractor="python-pptx") if markdown else None


def _xlsx(data: bytes, filename: str) -> Extracted | None:
    try:
        from openpyxl import load_workbook
    except Exception:  # noqa: BLE001
        return None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append("| " + " | ".join(cells) + " |")
            if rows:
                head, *rest = rows
                width = head.count("|") - 1
                sep = "|" + "|".join([" --- "] * width) + "|"
                blocks.append(f"## {sheet.title}\n\n" + "\n".join([head, sep, *rest]))
        markdown = _clean("\n\n".join(blocks))
    except Exception as exc:  # noqa: BLE001
        log.warning("openpyxl falhou em %s: %s", filename, exc)
        return None
    return Extracted(markdown=markdown, extractor="openpyxl") if markdown else None


# ── 3. texto puro ─────────────────────────────────────────────────────────


def _plain(data: bytes, filename: str) -> Extracted | None:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = data.decode("latin-1")
        except Exception:  # noqa: BLE001
            return None
    markdown = _clean(text)
    return Extracted(markdown=markdown, extractor="plain-text") if markdown else None


def _titulo(markdown: str, fallback: str) -> str:
    """Titulo do documento, com o frontmatter valendo mais que o corpo.

    Conceito OKF declara o nome em `title:`. Quando ele esta la, e melhor fonte
    do que o primeiro cabecalho: o corpo de um conceito costuma abrir por uma
    secao interna ("Definicao", "Calculo") que nao identifica coisa nenhuma numa
    lista de resultados.

    Isto vale INDEPENDENTE do modo OKF do Espaco. O modo decide o que e
    indexado; o titulo e so exibicao, e mostrar `---` nunca foi o comportamento
    pretendido em lugar nenhum.
    """
    concept = okf.parse(markdown)
    if concept and concept.title:
        return concept.title
    return _first_heading(markdown, fallback)


BY_EXTENSION = {
    ".pdf": (_docling, _pymupdf),
    ".docx": (_docling, _docx),
    ".doc": (_docling, _docx),
    ".pptx": (_docling, _pptx),
    ".ppt": (_docling, _pptx),
    ".xlsx": (_docling, _xlsx),
    ".xls": (_docling, _xlsx),
    ".html": (_docling, _plain),
    ".htm": (_docling, _plain),
}


def extract(data: bytes, filename: str) -> Extracted:
    """Documento canonico a partir do bruto. Levanta ExtractionError se nada der."""
    extension = Path(filename).suffix.lower()
    chain = BY_EXTENSION.get(extension)
    if chain is None:
        chain = (_plain,) if extension in PLAIN_EXTENSIONS else (_docling, _plain)

    tried: list[str] = []
    for extractor in chain:
        tried.append(extractor.__name__.strip("_"))
        result = extractor(data, filename)
        if result and result.markdown.strip():
            result.title = _titulo(result.markdown, Path(filename).stem)
            log.info("%s extraido por %s", filename, result.extractor)
            return result

    raise ExtractionError(
        f"nenhum extrator produziu texto para {filename} "
        f"(tentados: {', '.join(tried)}). "
        "Arquivo protegido por senha, corrompido ou so imagem sem OCR."
    )
